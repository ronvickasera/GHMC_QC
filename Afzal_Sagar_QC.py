import arcpy
import os
import re
from collections import Counter
from datetime import datetime, date

# ============================================================
# CONFIGURATION
# ============================================================

gdb_path = r"I:\PMU\NEOGEO COLONIES\AFZAL SAGAR\AFZAL SAGAR\AFZAL SAGAR.gdb"

ASSESSMENT_NAME = "Assesments"

# Put exact parcel feature class name here if known.
# Example:
# PARCEL_NAME = "Parcels"
#
# Keep None to automatically find a feature class
# containing "parcel".
PARCEL_NAME = None

output_excel = os.path.join(
    os.path.dirname(gdb_path),
    "AFZAL_SAGAR_QC_Report_With_OBJECTID.xlsx"
)


# ============================================================
# FIELD ALIASES
# ============================================================

FIELD_ALIASES = {

    "Building Type": [
        "Building_Type",
        "Building Type",
        "BuildingType",
        "BUILDINGTYPE",
        "BUILDING_TYPE"
    ],

    "Building ID": [
        "Building_ID",
        "Building ID",
        "BuildingID",
        "BUILDINGID",
        "BUILDING_ID"
    ],

    "Parcel ID": [
        "Parcel_ID",
        "Parcel ID",
        "ParcelID",
        "PARCELID",
        "PARCEL_ID"
    ],

    "Property ID": [
        "Property_ID",
        "Property ID",
        "PropertyID",
        "PROPERTYID",
        "PROPERTY_ID"
    ],

    "PTIN": [
        "PTIN",
        "P_TIN",
        "P-TIN",
        "Property_TIN",
        "PropertyTIN"
    ],

    "Owner Name": [
        "Owner_Name",
        "Owner Name",
        "OwnerName",
        "OWNERNAME",
        "OWNER_NAME"
    ],

    "Door Number": [
        "Door_Number",
        "Door Number",
        "DoorNumber",
        "DOORNUMBER",
        "DOOR_NUMBER"
    ],

    "Resident Owner Mobile Number": [
        "Resident_Owner_Mobile_Number",
        "Resident Owner Mobile Number",
        "ResidentOwnerMobileNumber",
        "Resident_Owner_Mobile",
        "Owner_Mobile",
        "Owner Mobile",
        "OwnerMobile",
        "Mobile_Number",
        "Mobile Number",
        "Mobile"
    ],

    "Photos": [
        "Photos",
        "Photo",
        "PHOTOS",
        "PHOTO",
        "Photo_Path",
        "Photo Path",
        "PhotoPath"
    ],

    "Floor Details": [
        "Floor_Details",
        "Floor Details",
        "FloorDetails",
        "FLOORDETAILS",
        "FLOOR_DETAILS",
        "Floor"
    ],

    "Existing Usage": [
        "Existing_Usage",
        "Existing Usage",
        "ExistingUsage",
        "EXISTINGUSAGE",
        "EXISTING_USAGE",
        "Existing_Use",
        "Existing Use"
    ],

    "Survey Start Date": [
        "Survey_Start_Date",
        "Survey Start Date",
        "SurveyStartDate",
        "SURVEY_START_DATE",
        "Start_Date",
        "Start Date",
        "SurveyDate"
    ],

    "Survey End Date": [
        "Survey_End_Date",
        "Survey End Date",
        "SurveyEndDate",
        "SURVEY_END_DATE",
        "End_Date",
        "End Date"
    ],

    "Total Plinth Area": [
        "Total_Plinth_Area",
        "Total Plinth Area",
        "TotalPlinthArea",
        "TOTALPLINTHAREA",
        "TOTAL_PLINTH_AREA",
        "Total_Plinth",
        "Total Plinth"
    ],

    "Existing Total Plinth Area": [
        "Existing_Total_Plinth_Area",
        "Existing Total Plinth Area",
        "ExistingTotalPlinthArea",
        "EXISTINGTOTALPLINTHAREA",
        "EXISTING_TOTAL_PLINTH_AREA",
        "Existing_Plinth_Area",
        "Existing Plinth Area",
        "ExistingTotalPlinth"
    ],

    "Plot Picture": [
        "Plot_Picture",
        "Plot Picture",
        "PlotPicture",
        "PLOTPICTURE",
        "PLOT_PICTURE",
        "Plot_Photo",
        "Plot Photo",
        "PlotPhoto"
    ]
}


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def normalize_name(value):

    return re.sub(
        r"[^a-z0-9]",
        "",
        str(value).lower()
    )


def find_field(fields, aliases):

    field_map = {}

    for field in fields:

        field_map[
            normalize_name(field.name)
        ] = field

        field_map[
            normalize_name(field.aliasName)
        ] = field

    for alias in aliases:

        key = normalize_name(alias)

        if key in field_map:

            return field_map[key]

    return None


def is_blank(value):

    if value is None:
        return True

    if isinstance(value, str):

        return value.strip() == ""

    return False


def validate_date(value):

    if value is None:
        return "BLANK"

    if isinstance(value, datetime):
        return "VALID"

    if isinstance(value, date):
        return "VALID"

    if isinstance(value, str):

        value = value.strip()

        if value == "":
            return "BLANK"

        formats = [
            "%Y-%m-%d",
            "%Y/%m/%d",
            "%d-%m-%Y",
            "%d/%m/%Y",
            "%m-%d-%Y",
            "%m/%d/%Y",
            "%Y-%m-%d %H:%M:%S",
            "%Y/%m/%d %H:%M:%S",
            "%d-%m-%Y %H:%M:%S",
            "%d/%m/%Y %H:%M:%S",
            "%m-%d-%Y %H:%M:%S",
            "%m/%d/%Y %H:%M:%S",
            "%d-%b-%Y",
            "%d-%B-%Y",
            "%d-%b-%Y %H:%M:%S",
            "%d-%B-%Y %H:%M:%S",
            "%Y%m%d"
        ]

        for fmt in formats:

            try:

                datetime.strptime(
                    value,
                    fmt
                )

                return "VALID"

            except ValueError:
                pass

        try:

            datetime.fromisoformat(value)

            return "VALID"

        except ValueError:

            return "INVALID"

    return "INVALID"


# ============================================================
# CHECK GDB
# ============================================================

print("=" * 100)
print("AFZAL SAGAR QC REPORT")
print("=" * 100)

if not arcpy.Exists(gdb_path):

    raise RuntimeError(
        "GDB does not exist:\n" +
        gdb_path
    )


# ============================================================
# FIND ASSESSMENTS
# ============================================================

assessment_fc = None

for root, dirs, files in arcpy.da.Walk(
    gdb_path,
    datatype="FeatureClass"
):

    for fc in files:

        if fc.lower() == ASSESSMENT_NAME.lower():

            assessment_fc = os.path.join(
                root,
                fc
            )

            break

    if assessment_fc:
        break


if assessment_fc is None:

    raise RuntimeError(
        "Assesments feature class was not found."
    )


print("\nAssessment Feature Class:")
print(assessment_fc)


# ============================================================
# FIND PARCELS
# ============================================================

parcel_fc = None

if PARCEL_NAME:

    for root, dirs, files in arcpy.da.Walk(
        gdb_path,
        datatype="FeatureClass"
    ):

        for fc in files:

            if fc.lower() == PARCEL_NAME.lower():

                parcel_fc = os.path.join(
                    root,
                    fc
                )

                break

        if parcel_fc:
            break

else:

    for root, dirs, files in arcpy.da.Walk(
        gdb_path,
        datatype="FeatureClass"
    ):

        for fc in files:

            if "parcel" in fc.lower():

                parcel_fc = os.path.join(
                    root,
                    fc
                )

                break

        if parcel_fc:
            break


if parcel_fc:

    print("\nParcel Feature Class:")
    print(parcel_fc)

else:

    print(
        "\nWARNING: Parcel feature class not found."
    )


# ============================================================
# TOTAL COUNTS
# ============================================================

total_assets = int(
    arcpy.management.GetCount(
        assessment_fc
    )[0]
)

if parcel_fc:

    total_parcels = int(
        arcpy.management.GetCount(
            parcel_fc
        )[0]
    )

else:

    total_parcels = 0


# ============================================================
# GET FIELDS
# ============================================================

fields = arcpy.ListFields(
    assessment_fc
)

data_fields = [
    f
    for f in fields
    if f.type not in [
        "Geometry",
        "Raster",
        "Blob"
    ]
]


# ============================================================
# RESOLVE FIELDS
# ============================================================

resolved = {}

for logical_name, aliases in FIELD_ALIASES.items():

    resolved[
        logical_name
    ] = find_field(
        data_fields,
        aliases
    )


# ============================================================
# OPENPYXL
# ============================================================

try:

    from openpyxl import Workbook
    from openpyxl.styles import (
        Font,
        PatternFill,
        Alignment
    )
    from openpyxl.utils import get_column_letter

except ImportError:

    raise RuntimeError(
        "openpyxl is not installed in ArcGIS Pro."
    )


# ============================================================
# CREATE WORKBOOK
# ============================================================

wb = Workbook()

ws_summary = wb.active
ws_summary.title = "QC Summary"

ws_categories = wb.create_sheet(
    "Category Counts"
)

ws_duplicates = wb.create_sheet(
    "Duplicate Check"
)

ws_blank = wb.create_sheet(
    "Blank Check"
)

ws_dates = wb.create_sheet(
    "Date Check"
)

ws_area = wb.create_sheet(
    "Area Summary"
)

ws_issues = wb.create_sheet(
    "QC Issues"
)

ws_mapping = wb.create_sheet(
    "Field Mapping"
)


# ============================================================
# STYLE
# ============================================================

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)


def format_header(ws):

    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


def resize_columns(ws):

    for column in ws.columns:

        max_length = 0

        letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            if cell.value is not None:

                length = len(
                    str(cell.value)
                )

                max_length = max(
                    max_length,
                    length
                )

        ws.column_dimensions[
            letter
        ].width = min(
            max_length + 2,
            60
        )


# ============================================================
# QC SUMMARY
# ============================================================

ws_summary.append([
    "QC Parameter",
    "Value",
    "Status"
])

format_header(ws_summary)

ws_summary.append([
    "Total Assets",
    total_assets,
    "INFO"
])

ws_summary.append([
    "Total Parcels",
    total_parcels,
    "INFO"
])


# ============================================================
# FIELD MAPPING
# ============================================================

ws_mapping.append([
    "Required Field",
    "Actual Field",
    "Type",
    "Status"
])

format_header(ws_mapping)

for logical_name, field in resolved.items():

    if field:

        ws_mapping.append([
            logical_name,
            field.name,
            field.type,
            "FOUND"
        ])

    else:

        ws_mapping.append([
            logical_name,
            "",
            "",
            "NOT FOUND"
        ])


# ============================================================
# CATEGORY COUNTS
# ============================================================

ws_categories.append([
    "Category Type",
    "Category",
    "Count",
    "Percentage"
])

format_header(ws_categories)


def write_categories(
    field,
    category_type
):

    if field is None:
        return

    counter = Counter()

    with arcpy.da.SearchCursor(
        assessment_fc,
        [field.name]
    ) as cursor:

        for row in cursor:

            value = row[0]

            if is_blank(value):

                value = "[BLANK]"

            else:

                value = str(
                    value
                ).strip()

            counter[value] += 1

    for category, count in sorted(
        counter.items(),
        key=lambda x: x[1],
        reverse=True
    ):

        percentage = 0

        if total_assets:

            percentage = (
                count /
                total_assets
            ) * 100

        ws_categories.append([
            category_type,
            category,
            count,
            round(
                percentage,
                2
            )
        ])


write_categories(
    resolved["Building Type"],
    "Building Type"
)

write_categories(
    resolved["Existing Usage"],
    "Existing Usage"
)


# ============================================================
# DUPLICATE CHECK WITH OBJECTID
# ============================================================

duplicate_fields = [

    "Building ID",
    "Parcel ID",
    "Property ID",
    "PTIN",
    "Owner Name",
    "Door Number",
    "Resident Owner Mobile Number"

]

ws_duplicates.append([
    "OBJECTID",
    "Field",
    "Actual Field",
    "Value",
    "Occurrence Count",
    "Status"
])

format_header(ws_duplicates)


for logical_name in duplicate_fields:

    field = resolved[
        logical_name
    ]

    if field is None:

        ws_duplicates.append([
            "",
            logical_name,
            "",
            "",
            "",
            "FIELD NOT FOUND"
        ])

        continue

    values = {}

    # --------------------------------------------------------
    # FIRST PASS - COLLECT VALUES + OBJECTIDS
    # --------------------------------------------------------

    with arcpy.da.SearchCursor(
        assessment_fc,
        ["OID@", field.name]
    ) as cursor:

        for row in cursor:

            oid = row[0]
            value = row[1]

            if is_blank(value):
                continue

            if isinstance(
                value,
                str
            ):

                value = value.strip()

            if value not in values:

                values[value] = []

            values[value].append(
                oid
            )

    # --------------------------------------------------------
    # WRITE DUPLICATES
    # --------------------------------------------------------

    duplicate_found = False

    for value, objectids in values.items():

        if len(objectids) > 1:

            duplicate_found = True

            for oid in objectids:

                ws_duplicates.append([
                    oid,
                    logical_name,
                    field.name,
                    value,
                    len(objectids),
                    "DUPLICATE"
                ])

    if not duplicate_found:

        ws_duplicates.append([
            "",
            logical_name,
            field.name,
            "",
            "",
            "PASS - NO DUPLICATES"
        ])


# ============================================================
# BLANK CHECK WITH OBJECTID
# ============================================================

blank_fields = [

    "Photos",
    "Floor Details",
    "Plot Picture"

]

ws_blank.append([
    "OBJECTID",
    "Field",
    "Actual Field",
    "Value",
    "Issue"
])

format_header(ws_blank)


for logical_name in blank_fields:

    field = resolved[
        logical_name
    ]

    if field is None:

        ws_blank.append([
            "",
            logical_name,
            "",
            "",
            "FIELD NOT FOUND"
        ])

        continue

    found_blank = False

    with arcpy.da.SearchCursor(
        assessment_fc,
        ["OID@", field.name]
    ) as cursor:

        for row in cursor:

            oid = row[0]
            value = row[1]

            if is_blank(value):

                found_blank = True

                ws_blank.append([
                    oid,
                    logical_name,
                    field.name,
                    "",
                    "BLANK / NULL"
                ])

    if not found_blank:

        ws_blank.append([
            "",
            logical_name,
            field.name,
            "",
            "PASS - NO BLANK VALUES"
        ])


# ============================================================
# DATE CHECK WITH OBJECTID
# ============================================================

date_fields = [

    "Survey Start Date",
    "Survey End Date"

]

ws_dates.append([
    "OBJECTID",
    "Field",
    "Actual Field",
    "Value",
    "Date Status"
])

format_header(ws_dates)


for logical_name in date_fields:

    field = resolved[
        logical_name
    ]

    if field is None:

        ws_dates.append([
            "",
            logical_name,
            "",
            "",
            "FIELD NOT FOUND"
        ])

        continue

    found_issue = False

    with arcpy.da.SearchCursor(
        assessment_fc,
        ["OID@", field.name]
    ) as cursor:

        for row in cursor:

            oid = row[0]
            value = row[1]

            result = validate_date(
                value
            )

            if result == "INVALID":

                found_issue = True

                ws_dates.append([
                    oid,
                    logical_name,
                    field.name,
                    value,
                    "INVALID DATE FORMAT"
                ])

            elif result == "BLANK":

                found_issue = True

                ws_dates.append([
                    oid,
                    logical_name,
                    field.name,
                    "",
                    "BLANK DATE"
                ])

            elif result == "VALID":

                if isinstance(
                    value,
                    (datetime, date)
                ):

                    if value > datetime.now():

                        found_issue = True

                        ws_dates.append([
                            oid,
                            logical_name,
                            field.name,
                            value,
                            "FUTURE DATE"
                        ])

    if not found_issue:

        ws_dates.append([
            "",
            logical_name,
            field.name,
            "",
            "PASS"
        ])


# ============================================================
# AREA CALCULATION
# ============================================================

total_plinth_field = resolved[
    "Total Plinth Area"
]

existing_plinth_field = resolved[
    "Existing Total Plinth Area"
]

total_plinth = 0
existing_plinth = 0

total_plinth_blank = 0
existing_plinth_blank = 0


if total_plinth_field:

    with arcpy.da.SearchCursor(
        assessment_fc,
        [total_plinth_field.name]
    ) as cursor:

        for row in cursor:

            value = row[0]

            if is_blank(value):

                total_plinth_blank += 1

                continue

            try:

                total_plinth += float(
                    value
                )

            except:

                pass


if existing_plinth_field:

    with arcpy.da.SearchCursor(
        assessment_fc,
        [existing_plinth_field.name]
    ) as cursor:

        for row in cursor:

            value = row[0]

            if is_blank(value):

                existing_plinth_blank += 1

                continue

            try:

                existing_plinth += float(
                    value
                )

            except:

                pass


plinth_difference = (
    total_plinth -
    existing_plinth
)


# ============================================================
# AREA SUMMARY
# ============================================================

ws_area.append([
    "Parameter",
    "Field",
    "Value",
    "Blank / NULL"
])

format_header(ws_area)

ws_area.append([
    "Total Plinth Area",
    total_plinth_field.name
    if total_plinth_field
    else "NOT FOUND",
    total_plinth,
    total_plinth_blank
])

ws_area.append([
    "Existing Total Plinth Area",
    existing_plinth_field.name
    if existing_plinth_field
    else "NOT FOUND",
    existing_plinth,
    existing_plinth_blank
])

ws_area.append([
    "Difference",
    "Total Plinth - Existing Total Plinth",
    plinth_difference,
    ""
])


# ============================================================
# QC SUMMARY AREA
# ============================================================

ws_summary.append([
    "Total Plinth Area",
    total_plinth,
    "INFO"
])

ws_summary.append([
    "Existing Total Plinth Area",
    existing_plinth,
    "INFO"
])

ws_summary.append([
    "Plinth Area Difference",
    plinth_difference,
    "INFO"
])


# ============================================================
# QC ISSUES
# ============================================================

ws_issues.append([
    "OBJECTID",
    "Field",
    "Value",
    "Issue"
])

format_header(ws_issues)


# Add issues from blank sheet
for row in ws_blank.iter_rows(
    min_row=2,
    values_only=True
):

    oid = row[0]

    if oid not in ("", None):

        ws_issues.append([
            oid,
            row[1],
            row[3],
            row[4]
        ])


# Add issues from date sheet
for row in ws_dates.iter_rows(
    min_row=2,
    values_only=True
):

    oid = row[0]

    if oid not in ("", None):

        ws_issues.append([
            oid,
            row[1],
            row[3],
            row[4]
        ])


# ============================================================
# FORMAT ALL SHEETS
# ============================================================

for ws in wb.worksheets:

    ws.freeze_panes = "A2"

    resize_columns(ws)


# ============================================================
# SAVE
# ============================================================

print("\nSaving Excel report...")

wb.save(
    output_excel
)

print("\n")
print("=" * 100)
print("QC REPORT GENERATED")
print("=" * 100)

print("\nTotal Assets:", total_assets)

print("Total Parcels:", total_parcels)

print(
    "Total Plinth Area:",
    total_plinth
)

print(
    "Existing Total Plinth Area:",
    existing_plinth
)

print(
    "Difference:",
    plinth_difference
)

print("\nOutput:")
print(output_excel)

print("\nCompleted.")
print("=" * 100)