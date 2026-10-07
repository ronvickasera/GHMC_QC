import arcpy
import os
from collections import Counter
from datetime import datetime, date

# ============================================================
# 1. CONFIGURATION
# ============================================================

GDB_PATH = r"I:\PMU\NEOGEO COLONIES\ADARSH NAGAR\ADARSH NAGAR\ADARSH NAGAR.gdb"

FEATURE_CLASS_NAME = "Assesments"

OUTPUT_FOLDER = r"I:\PMU\NEOGEO COLONIES\ADARSH NAGAR\ADARSH NAGAR"

OUTPUT_EXCEL = os.path.join(
    OUTPUT_FOLDER,
    "Assesments_Data_Quality_Report.xlsx"
)

MAX_DUPLICATES = 100
MAX_SAMPLE_VALUES = 20


# ============================================================
# 2. CHECK GDB
# ============================================================

print("=" * 100)
print("ASSESMENTS DATA QUALITY REPORT")
print("=" * 100)

if not arcpy.Exists(GDB_PATH):
    raise RuntimeError(
        "Geodatabase not found:\n" + GDB_PATH
    )


# ============================================================
# 3. FIND ASSESMENTS FEATURE CLASS
# ============================================================

print("\nSearching for Assesments...")

found_fc = None

for root, dirs, files in arcpy.da.Walk(
    GDB_PATH,
    datatype="FeatureClass"
):

    for fc in files:

        if fc.lower() == FEATURE_CLASS_NAME.lower():

            found_fc = os.path.join(root, fc)

            break

    if found_fc:
        break


if found_fc is None:

    raise RuntimeError(
        "Feature class '" +
        FEATURE_CLASS_NAME +
        "' was not found."
    )


print("\nFeature class found:")
print(found_fc)


# ============================================================
# 4. IMPORT OPENPYXL
# ============================================================

try:

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

except ImportError:

    raise RuntimeError(
        "openpyxl is not installed in ArcGIS Pro Python."
    )


# ============================================================
# 5. GET FIELDS
# ============================================================

all_fields = arcpy.ListFields(found_fc)

fields = []

for f in all_fields:

    if f.type not in ["Geometry", "Raster", "Blob"]:

        fields.append(f)


print("\nTotal fields:", len(fields))


# ============================================================
# 6. RECORD COUNT
# ============================================================

total_records = int(
    arcpy.management.GetCount(found_fc)[0]
)

print("Total records:", total_records)


# ============================================================
# 7. DATE VALIDATION FUNCTION
# ============================================================

DATE_FORMATS = [

    "%Y-%m-%d",
    "%Y/%m/%d",

    "%d-%m-%Y",
    "%d/%m/%Y",

    "%m-%d-%Y",
    "%m/%d/%Y",

    "%Y-%m-%d %H:%M:%S",
    "%Y/%m/%d %H:%M:%S",

    "%d-%m-%Y %H:%M:%S",
    "%d/%m/%Y %H:%M:%S",

    "%m-%d-%Y %H:%M:%S",
    "%m/%d/%Y %H:%M:%S",

    "%d-%b-%Y",
    "%d-%B-%Y",

    "%d-%b-%Y %H:%M:%S",
    "%d-%B-%Y %H:%M:%S",

    "%Y%m%d"
]


def check_date(value):

    if value is None:

        return None

    if isinstance(value, datetime):

        return True

    if isinstance(value, date):

        return True

    if isinstance(value, str):

        value = value.strip()

        if value == "":

            return None

        for fmt in DATE_FORMATS:

            try:

                datetime.strptime(
                    value,
                    fmt
                )

                return True

            except ValueError:

                continue

        try:

            datetime.fromisoformat(value)

            return True

        except ValueError:

            return False

    return False


# ============================================================
# 8. CREATE EXCEL WORKBOOK
# ============================================================

wb = Workbook()

ws_summary = wb.active
ws_summary.title = "Field Summary"

ws_duplicates = wb.create_sheet(
    "Duplicate Values"
)

ws_dates = wb.create_sheet(
    "Invalid Dates"
)

ws_samples = wb.create_sheet(
    "Sample Values"
)

ws_info = wb.create_sheet(
    "Record Summary"
)


# ============================================================
# 9. EXCEL STYLING
# ============================================================

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)


def style_header(ws):

    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center"
        )


def resize_columns(ws):

    for column in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column[0].column
        )

        for cell in column:

            if cell.value is not None:

                length = len(
                    str(cell.value)
                )

                if length > max_length:

                    max_length = length

        if max_length > 60:

            max_length = 60

        ws.column_dimensions[
            column_letter
        ].width = max_length + 2


# ============================================================
# 10. SUMMARY HEADERS
# ============================================================

summary_headers = [

    "Field Name",
    "Alias",
    "Data Type",
    "Length",
    "Total Records",
    "NULL Count",
    "NULL %",
    "Blank Count",
    "Blank %",
    "Non-Null Count",
    "Unique Value Count",
    "Duplicate Value Count",
    "Duplicate Record Count",
    "Has Duplicate",
    "Has NULL",
    "Has Blank",
    "Date Check",
    "Invalid Date Count",
    "Future Date Count"

]

ws_summary.append(summary_headers)

style_header(ws_summary)


# ============================================================
# 11. DUPLICATE HEADERS
# ============================================================

ws_duplicates.append([

    "Field Name",
    "Field Type",
    "Value",
    "Count"

])

style_header(ws_duplicates)


# ============================================================
# 12. DATE HEADERS
# ============================================================

ws_dates.append([

    "Field Name",
    "OBJECTID",
    "Value",
    "Issue"

])

style_header(ws_dates)


# ============================================================
# 13. SAMPLE HEADERS
# ============================================================

ws_samples.append([

    "Field Name",
    "Field Type",
    "Sample Value"

])

style_header(ws_samples)


# ============================================================
# 14. GENERAL INFORMATION
# ============================================================

ws_info.append([
    "Property",
    "Value"
])

style_header(ws_info)

ws_info.append([
    "Geodatabase",
    GDB_PATH
])

ws_info.append([
    "Feature Class",
    found_fc
])

ws_info.append([
    "Total Records",
    total_records
])

ws_info.append([
    "Total Fields",
    len(fields)
])

ws_info.append([
    "Report Generated",
    datetime.now()
])


# ============================================================
# 15. PROCESS EACH FIELD
# ============================================================

for index, field in enumerate(fields, 1):

    field_name = field.name
    field_type = field.type
    field_alias = field.aliasName
    field_length = field.length

    print(
        "\n[" +
        str(index) +
        "/" +
        str(len(fields)) +
        "] Checking: " +
        field_name
    )

    null_count = 0
    blank_count = 0

    counter = Counter()

    samples = []

    invalid_date_count = 0
    future_date_count = 0

    # --------------------------------------------------------
    # READ DATA
    # --------------------------------------------------------

    with arcpy.da.SearchCursor(
        found_fc,
        ["OID@", field_name]
    ) as cursor:

        for row in cursor:

            oid = row[0]
            value = row[1]

            # ------------------------------------------------
            # NULL
            # ------------------------------------------------

            if value is None:

                null_count += 1

                continue

            # ------------------------------------------------
            # BLANK
            # ------------------------------------------------

            if isinstance(value, str):

                cleaned_value = value.strip()

                if cleaned_value == "":

                    blank_count += 1

                value_to_count = cleaned_value

            else:

                value_to_count = value

            # ------------------------------------------------
            # COUNT VALUES
            # ------------------------------------------------

            try:

                counter[value_to_count] += 1

            except TypeError:

                counter[
                    str(value_to_count)
                ] += 1

            # ------------------------------------------------
            # SAMPLE VALUES
            # ------------------------------------------------

            if len(samples) < MAX_SAMPLE_VALUES:

                if value_to_count not in samples:

                    samples.append(
                        value_to_count
                    )

            # ------------------------------------------------
            # DATE FIELD
            # ------------------------------------------------

            if field_type == "Date":

                result = check_date(value)

                if result is False:

                    invalid_date_count += 1

                    ws_dates.append([

                        field_name,
                        oid,
                        value,
                        "Invalid Date"

                    ])

                elif result is True:

                    if isinstance(
                        value,
                        datetime
                    ):

                        if value > datetime.now():

                            future_date_count += 1

                            ws_dates.append([

                                field_name,
                                oid,
                                value,
                                "Future Date"

                            ])

            # ------------------------------------------------
            # DATE-LIKE TEXT FIELD
            # ------------------------------------------------

            elif field_type in [
                "String",
                "Text"
            ]:

                lower_name = field_name.lower()

                date_keywords = [

                    "date",
                    "dob",
                    "birth",
                    "surveydate",
                    "assessmentdate",
                    "createddate",
                    "updateddate",
                    "completiondate"

                ]

                looks_like_date = False

                for keyword in date_keywords:

                    if keyword in lower_name:

                        looks_like_date = True

                        break

                if looks_like_date:

                    result = check_date(value)

                    if result is False:

                        invalid_date_count += 1

                        ws_dates.append([

                            field_name,
                            oid,
                            value,
                            "Invalid Date Format"

                        ])

    # ========================================================
    # DUPLICATE ANALYSIS
    # ========================================================

    duplicate_values = {}

    for value, count in counter.items():

        if count > 1:

            duplicate_values[
                value
            ] = count

    duplicate_value_count = len(
        duplicate_values
    )

    duplicate_record_count = sum(
        duplicate_values.values()
    )

    # ========================================================
    # WRITE DUPLICATES
    # ========================================================

    sorted_duplicates = sorted(

        duplicate_values.items(),

        key=lambda item: item[1],

        reverse=True

    )

    for value, count in sorted_duplicates[
        :MAX_DUPLICATES
    ]:

        ws_duplicates.append([

            field_name,
            field_type,
            value,
            count

        ])

    # ========================================================
    # WRITE SAMPLE VALUES
    # ========================================================

    for sample in samples:

        ws_samples.append([

            field_name,
            field_type,
            sample

        ])

    # ========================================================
    # PERCENTAGES
    # ========================================================

    if total_records > 0:

        null_percentage = (
            null_count /
            total_records
        ) * 100

        blank_percentage = (
            blank_count /
            total_records
        ) * 100

    else:

        null_percentage = 0
        blank_percentage = 0

    # ========================================================
    # DATE STATUS
    # ========================================================

    if field_type == "Date":

        if invalid_date_count == 0:

            date_status = "PASS"

        else:

            date_status = "FAIL"

    elif invalid_date_count > 0:

        date_status = (
            "DATE-LIKE TEXT - INVALID"
        )

    else:

        date_status = "N/A"

    # ========================================================
    # WRITE SUMMARY
    # ========================================================

    ws_summary.append([

        field_name,
        field_alias,
        field_type,
        field_length,

        total_records,

        null_count,
        round(null_percentage, 2),

        blank_count,
        round(blank_percentage, 2),

        total_records - null_count,

        len(counter),

        duplicate_value_count,

        duplicate_record_count,

        "YES"
        if duplicate_value_count > 0
        else "NO",

        "YES"
        if null_count > 0
        else "NO",

        "YES"
        if blank_count > 0
        else "NO",

        date_status,

        invalid_date_count,

        future_date_count

    ])


# ============================================================
# 16. OVERALL SUMMARY
# ============================================================

fields_with_null = 0
fields_with_blank = 0
fields_with_duplicates = 0
fields_with_invalid_dates = 0

for row in ws_summary.iter_rows(
    min_row=2,
    values_only=True
):

    if row[5] > 0:

        fields_with_null += 1

    if row[7] > 0:

        fields_with_blank += 1

    if row[11] > 0:

        fields_with_duplicates += 1

    if row[17] > 0:

        fields_with_invalid_dates += 1


ws_info.append([
    "Fields with NULL values",
    fields_with_null
])

ws_info.append([
    "Fields with blank values",
    fields_with_blank
])

ws_info.append([
    "Fields with duplicate values",
    fields_with_duplicates
])

ws_info.append([
    "Fields with invalid dates",
    fields_with_invalid_dates
])


# ============================================================
# 17. FORMAT WORKSHEETS
# ============================================================

for ws in wb.worksheets:

    ws.freeze_panes = "A2"

    resize_columns(ws)


# ============================================================
# 18. SAVE EXCEL
# ============================================================

print("\n")
print("=" * 100)
print("Saving Excel report...")
print("=" * 100)

wb.save(OUTPUT_EXCEL)

print("\nSUCCESS!")
print("------------------------------------------------------------")
print("Feature Class:")
print(found_fc)

print("\nExcel Report:")
print(OUTPUT_EXCEL)

print("\nTotal Records:")
print(total_records)

print("\nTotal Fields:")
print(len(fields))

print("\nReport completed successfully.")
print("=" * 100)