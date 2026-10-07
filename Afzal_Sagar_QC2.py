import arcpy
import os
from collections import defaultdict, Counter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

gdb_path = r"I:\PMU\NEOGEO COLONIES\AFZAL SAGAR\AFZAL SAGAR\AFZAL SAGAR.gdb"

ASSESSMENT_NAME = "Assesments"

output_excel = os.path.join(
    os.path.dirname(gdb_path),
    "AFZAL_SAGAR_ASSESSMENT_QC_REPORT.xlsx"
)


# ============================================================
# FIELD ALIASES
# ============================================================

FIELD_ALIASES = {

    "PTIN": [
        "PTIN",
        "P_TIN",
        "P-TIN",
        "Property_TIN",
        "PropertyTIN"
    ],

    "Owner Name": [
        "Owner_Name",
        "Owner Name",
        "OwnerName",
        "OWNERNAME",
        "OWNER_NAME"
    ],

    "Door Number": [
        "Door_Number",
        "Door Number",
        "DoorNumber",
        "DOORNUMBER",
        "DOOR_NUMBER"
    ],

    "Phone Number": [
        "Residential_Owner_Mobile_Number",
        "Residential Owner Mobile Number",
        "ResidentialOwnerMobileNumber",

        "Resident_Owner_Mobile_Number",
        "Resident Owner Mobile Number",
        "ResidentOwnerMobileNumber",

        "Resident_Owner_Mobile",
        "Owner_Mobile",
        "Owner Mobile",
        "OwnerMobile",

        "Mobile_Number",
        "Mobile Number",
        "Mobile",
        "Phone_Number",
        "Phone Number",
        "PhoneNumber",
        "Phone"
    ],

    "Existing Usage": [
        "Existing_Usage",
        "Existing Usage",
        "ExistingUsage",
        "EXISTINGUSAGE",
        "EXISTING_USAGE",
        "Existing_Use",
        "Existing Use"
    ]
}


# ============================================================
# N/A VALUES
# ============================================================

NA_VALUES = {
    "",
    "N/A",
    "NA",
    "N.A",
    "N.A.",
    "NULL",
    "NONE",
    "NOT AVAILABLE",
    "NOTAVAILABLE",
    "-",
    "--"
}


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def normalize_value(value):
    """
    Converts a field value into a normalized string
    for duplicate/N/A checking.
    """

    if value is None:
        return ""

    return str(value).strip()


def is_na_value(value):
    """
    Checks NULL, blank and common N/A representations.
    """

    if value is None:
        return True

    value = str(value).strip()

    if value == "":
        return True

    return value.upper() in NA_VALUES


def normalize_field_name(value):

    if value is None:
        return ""

    return "".join(
        ch.lower()
        for ch in str(value)
        if ch.isalnum()
    )


def find_field(fields, aliases):

    field_map = {}

    for field in fields:

        field_map[
            normalize_field_name(field.name)
        ] = field

        field_map[
            normalize_field_name(field.aliasName)
        ] = field

    for alias in aliases:

        key = normalize_field_name(alias)

        if key in field_map:
            return field_map[key]

    return None


def find_assessment_fc(gdb, assessment_name):

    for root, dirs, files in arcpy.da.Walk(
        gdb,
        datatype="FeatureClass"
    ):

        for fc in files:

            if fc.lower() == assessment_name.lower():

                return os.path.join(root, fc)

    return None


# ============================================================
# CHECK GDB
# ============================================================

print("=" * 100)
print("ASSESSMENT QC")
print("=" * 100)

if not arcpy.Exists(gdb_path):

    raise RuntimeError(
        "GDB does not exist:\n{}".format(gdb_path)
    )


# ============================================================
# FIND ASSESSMENT FEATURE CLASS
# ============================================================

assessment_fc = find_assessment_fc(
    gdb_path,
    ASSESSMENT_NAME
)

if assessment_fc is None:

    raise RuntimeError(
        "Assessment feature class '{}' was not found.".format(
            ASSESSMENT_NAME
        )
    )


print("\nAssessment Feature Class:")
print(assessment_fc)


# ============================================================
# TOTAL PROPERTIES
# ============================================================

total_properties = int(
    arcpy.management.GetCount(
        assessment_fc
    )[0]
)

print("\nTotal Properties:")
print(total_properties)


# ============================================================
# GET FIELDS
# ============================================================

fields = arcpy.ListFields(
    assessment_fc
)

data_fields = [
    f
    for f in fields
    if f.type not in [
        "Geometry",
        "Raster",
        "Blob"
    ]
]


# ============================================================
# RESOLVE FIELDS
# ============================================================

resolved = {}

for logical_name, aliases in FIELD_ALIASES.items():

    resolved[logical_name] = find_field(
        data_fields,
        aliases
    )


print("\n")
print("=" * 100)
print("FIELD MAPPING")
print("=" * 100)

for logical_name, field in resolved.items():

    if field:

        print(
            "{:<25} -> {}".format(
                logical_name,
                field.name
            )
        )

    else:

        print(
            "{:<25} -> NOT FOUND".format(
                logical_name
            )
        )


# ============================================================
# CHECK REQUIRED FIELDS
# ============================================================

required_fields = [
    "PTIN",
    "Owner Name",
    "Door Number",
    "Phone Number",
    "Existing Usage"
]

missing_fields = []

for field_name in required_fields:

    if resolved[field_name] is None:

        missing_fields.append(field_name)


if missing_fields:

    print("\nWARNING - Missing fields:")

    for field_name in missing_fields:
        print("  -", field_name)


# ============================================================
# CREATE EXCEL WORKBOOK
# ============================================================

wb = Workbook()

ws_summary = wb.active
ws_summary.title = "QC Summary"

ws_ptin = wb.create_sheet(
    "PTIN Duplicate"
)

ws_owner = wb.create_sheet(
    "Owner Duplicate"
)

ws_door = wb.create_sheet(
    "Door Duplicate"
)

ws_phone = wb.create_sheet(
    "Phone Duplicate"
)

ws_na = wb.create_sheet(
    "N-A Values"
)

ws_usage = wb.create_sheet(
    "Existing Usage"
)

ws_errors = wb.create_sheet(
    "Record Error Details"
)

ws_mapping = wb.create_sheet(
    "Field Mapping"
)


# ============================================================
# EXCEL STYLING
# ============================================================

header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

thin_side = Side(
    style="thin",
    color="B7B7B7"
)

thin_border = Border(
    left=thin_side,
    right=thin_side,
    top=thin_side,
    bottom=thin_side
)


def format_header(ws):

    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        cell.border = thin_border


def style_all_cells(ws):

    for row in ws.iter_rows():

        for cell in row:

            cell.border = thin_border

            cell.alignment = Alignment(
                vertical="top"
            )


def resize_columns(ws):

    for column in ws.columns:

        if not column:
            continue

        max_length = 0

        for cell in column:

            if cell.value is not None:

                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        letter = get_column_letter(
            column[0].column
        )

        ws.column_dimensions[
            letter
        ].width = min(
            max_length + 2,
            60
        )


# ============================================================
# FIELD MAPPING SHEET
# ============================================================

ws_mapping.append([
    "Logical Field",
    "Actual Field",
    "Field Type",
    "Status"
])

format_header(ws_mapping)

for logical_name, field in resolved.items():

    if field:

        ws_mapping.append([
            logical_name,
            field.name,
            field.type,
            "FOUND"
        ])

    else:

        ws_mapping.append([
            logical_name,
            "",
            "",
            "NOT FOUND"
        ])


# ============================================================
# RECORD ERROR TRACKING
# ============================================================

record_errors = defaultdict(list)


def add_record_error(
    oid,
    error_type,
    field_name,
    value
):

    record_errors[oid].append({
        "error_type": error_type,
        "field": field_name,
        "value": value
    })


# ============================================================
# GENERIC DUPLICATE CHECK
# ============================================================

def duplicate_check(
    logical_name,
    worksheet,
    error_type
):

    field = resolved.get(logical_name)

    worksheet.append([
        "OBJECTID",
        "Field",
        "Value",
        "Occurrence Count",
        "Status"
    ])

    format_header(worksheet)

    if field is None:

        worksheet.append([
            "",
            logical_name,
            "",
            "",
            "FIELD NOT FOUND"
        ])

        return {
            "duplicate_values": 0,
            "duplicate_records": 0
        }


    value_objects = defaultdict(list)


    # --------------------------------------------------------
    # READ VALUES
    # --------------------------------------------------------

    with arcpy.da.SearchCursor(
        assessment_fc,
        [
            "OID@",
            field.name
        ]
    ) as cursor:

        for oid, value in cursor:

            if is_na_value(value):
                continue

            cleaned_value = normalize_value(value)

            value_objects[
                cleaned_value
            ].append(oid)


    duplicate_values = 0
    duplicate_records = 0


    # --------------------------------------------------------
    # WRITE DUPLICATES
    # --------------------------------------------------------

    for value, objectids in sorted(
        value_objects.items(),
        key=lambda x: str(x[0])
    ):

        occurrence_count = len(
            objectids
        )

        if occurrence_count > 1:

            duplicate_values += 1

            duplicate_records += occurrence_count

            for oid in objectids:

                worksheet.append([
                    oid,
                    field.name,
                    value,
                    occurrence_count,
                    "DUPLICATE"
                ])

                add_record_error(
                    oid,
                    error_type,
                    field.name,
                    value
                )


    if duplicate_values == 0:

        worksheet.append([
            "",
            field.name,
            "",
            "",
            "PASS - NO DUPLICATES"
        ])


    return {
        "duplicate_values": duplicate_values,
        "duplicate_records": duplicate_records
    }


# ============================================================
# PTIN DUPLICATE CHECK
# ============================================================

print("\nChecking PTIN duplicates...")

ptin_result = duplicate_check(
    "PTIN",
    ws_ptin,
    "DUPLICATE PTIN"
)


# ============================================================
# OWNER NAME DUPLICATE CHECK
# ============================================================

print("Checking Owner Name duplicates...")

owner_result = duplicate_check(
    "Owner Name",
    ws_owner,
    "DUPLICATE OWNER NAME"
)


# ============================================================
# DOOR NUMBER DUPLICATE CHECK
# ============================================================

print("Checking Door Number duplicates...")

door_result = duplicate_check(
    "Door Number",
    ws_door,
    "DUPLICATE DOOR NUMBER"
)


# ============================================================
# PHONE NUMBER DUPLICATE CHECK
# ============================================================

print("Checking Phone Number duplicates...")

phone_result = duplicate_check(
    "Phone Number",
    ws_phone,
    "DUPLICATE PHONE NUMBER"
)


# ============================================================
# N/A VALUE CHECK
# ============================================================

ws_na.append([
    "OBJECTID",
    "Field",
    "Value",
    "Issue"
])

format_header(ws_na)


na_statistics = {}


def check_na_values(
    logical_name
):

    field = resolved.get(
        logical_name
    )

    if field is None:

        ws_na.append([
            "",
            logical_name,
            "",
            "FIELD NOT FOUND"
        ])

        return 0


    na_count = 0


    with arcpy.da.SearchCursor(
        assessment_fc,
        [
            "OID@",
            field.name
        ]
    ) as cursor:

        for oid, value in cursor:

            if is_na_value(value):

                na_count += 1

                display_value = (
                    ""
                    if value is None
                    else str(value)
                )

                ws_na.append([
                    oid,
                    field.name,
                    display_value,
                    "N/A / NULL / BLANK"
                ])

                add_record_error(
                    oid,
                    "N/A VALUE",
                    field.name,
                    display_value
                )


    return na_count


# Owner Name N/A
owner_na = check_na_values(
    "Owner Name"
)

na_statistics[
    "Owner Name"
] = owner_na


# Door Number N/A
door_na = check_na_values(
    "Door Number"
)

na_statistics[
    "Door Number"
] = door_na


# Phone Number N/A
phone_na = check_na_values(
    "Phone Number"
)

na_statistics[
    "Phone Number"
] = phone_na


# ============================================================
# EXISTING USAGE CATEGORY
# ============================================================

print("\nChecking Existing Usage categories...")

ws_usage.append([
    "Existing Usage",
    "Count",
    "Percentage"
])

format_header(ws_usage)


usage_field = resolved.get(
    "Existing Usage"
)

usage_counter = Counter()


if usage_field:

    with arcpy.da.SearchCursor(
        assessment_fc,
        [
            usage_field.name
        ]
    ) as cursor:

        for row in cursor:

            value = row[0]

            if is_na_value(value):

                value = "[N/A / BLANK]"

            else:

                value = normalize_value(
                    value
                )

            usage_counter[
                value
            ] += 1


    for category, count in sorted(
        usage_counter.items(),
        key=lambda x: x[1],
        reverse=True
    ):

        percentage = (
            count / total_properties * 100
            if total_properties
            else 0
        )

        ws_usage.append([
            category,
            count,
            round(
                percentage,
                2
            )
        ])

else:

    ws_usage.append([
        "[FIELD NOT FOUND]",
        "",
        ""
    ])


# ============================================================
# RECORD ERROR DETAILS
# ============================================================

ws_errors.append([
    "OBJECTID",
    "Error Count",
    "Error Percentage",
    "QC Status",
    "Error Types",
    "Fields"
])

format_header(ws_errors)


# ============================================================
# ADD QC FIELDS TO ASSESSMENT FEATURE CLASS
# ============================================================

existing_field_names = [
    f.name
    for f in arcpy.ListFields(
        assessment_fc
    )
]


def add_field_if_missing(
    field_name,
    field_type
):

    if field_name not in existing_field_names:

        arcpy.management.AddField(
            assessment_fc,
            field_name,
            field_type
        )

        print(
            "Added field:",
            field_name
        )


add_field_if_missing(
    "QC_Error_Count",
    "LONG"
)

add_field_if_missing(
    "QC_Error_Pct",
    "DOUBLE"
)

add_field_if_missing(
    "QC_Status",
    "TEXT"
)


# ============================================================
# UPDATE ASSESSMENT RECORDS
# ============================================================

print("\nUpdating Assessment QC fields...")


with arcpy.da.UpdateCursor(
    assessment_fc,
    [
        "OID@",
        "QC_Error_Count",
        "QC_Error_Pct",
        "QC_Status"
    ]
) as cursor:

    for row in cursor:

        oid = row[0]

        errors = record_errors.get(
            oid,
            []
        )

        error_count = len(errors)

        error_percentage = (
            error_count / len(required_fields) * 100
        )

        if error_count > 0:

            status = "ERROR"

        else:

            status = "PASS"


        row[1] = error_count

        row[2] = round(
            error_percentage,
            2
        )

        row[3] = status

        cursor.updateRow(row)


        # ----------------------------------------------------
        # EXCEL RECORD DETAILS
        # ----------------------------------------------------

        if errors:

            error_types = "; ".join(
                sorted(
                    set(
                        e["error_type"]
                        for e in errors
                    )
                )
            )

            fields_with_errors = "; ".join(
                sorted(
                    set(
                        e["field"]
                        for e in errors
                    )
                )
            )

            ws_errors.append([
                oid,
                error_count,
                round(
                    error_percentage,
                    2
                ),
                status,
                error_types,
                fields_with_errors
            ])


# ============================================================
# CALCULATE OVERALL ERROR
# ============================================================

error_records = len(
    record_errors
)

clean_records = (
    total_properties -
    error_records
)

overall_error_percentage = (
    error_records /
    total_properties *
    100
    if total_properties
    else 0
)

overall_pass_percentage = (
    clean_records /
    total_properties *
    100
    if total_properties
    else 0
)


# ============================================================
# FIELD LEVEL ERROR PERCENTAGES
# ============================================================

ptin_error_pct = (
    ptin_result["duplicate_records"]
    / total_properties * 100
    if total_properties
    else 0
)

owner_error_count = (
    owner_result["duplicate_records"]
    + owner_na
)

owner_error_pct = (
    owner_error_count /
    total_properties * 100
    if total_properties
    else 0
)

door_error_count = (
    door_result["duplicate_records"]
    + door_na
)

door_error_pct = (
    door_error_count /
    total_properties * 100
    if total_properties
    else 0
)

phone_error_count = (
    phone_result["duplicate_records"]
    + phone_na
)

phone_error_pct = (
    phone_error_count /
    total_properties * 100
    if total_properties
    else 0
)


# ============================================================
# QC SUMMARY
# ============================================================

ws_summary.append([
    "QC Parameter",
    "Total Properties",
    "Error Count",
    "Error Percentage",
    "Status"
])

format_header(ws_summary)


# ------------------------------------------------------------
# TOTAL PROPERTIES
# ------------------------------------------------------------

ws_summary.append([
    "Total Properties",
    total_properties,
    "",
    "",
    "INFO"
])


# ------------------------------------------------------------
# PTIN
# ------------------------------------------------------------

ws_summary.append([
    "PTIN - Duplicate Records",
    total_properties,
    ptin_result[
        "duplicate_records"
    ],
    round(
        ptin_error_pct,
        2
    ),
    (
        "PASS"
        if ptin_result[
            "duplicate_records"
        ] == 0
        else "ERROR"
    )
])


# ------------------------------------------------------------
# OWNER
# ------------------------------------------------------------

ws_summary.append([
    "Owner Name - Duplicate Records",
    total_properties,
    owner_result[
        "duplicate_records"
    ],
    round(
        (
            owner_result[
                "duplicate_records"
            ] /
            total_properties *
            100
        )
        if total_properties
        else 0,
        2
    ),
    (
        "PASS"
        if owner_result[
            "duplicate_records"
        ] == 0
        else "ERROR"
    )
])


ws_summary.append([
    "Owner Name - N/A",
    total_properties,
    owner_na,
    round(
        owner_na /
        total_properties *
        100
        if total_properties
        else 0,
        2
    ),
    (
        "PASS"
        if owner_na == 0
        else "ERROR"
    )
])


# ------------------------------------------------------------
# DOOR
# ------------------------------------------------------------

ws_summary.append([
    "Door Number - Duplicate Records",
    total_properties,
    door_result[
        "duplicate_records"
    ],
    round(
        (
            door_result[
                "duplicate_records"
            ] /
            total_properties *
            100
        )
        if total_properties
        else 0,
        2
    ),
    (
        "PASS"
        if door_result[
            "duplicate_records"
        ] == 0
        else "ERROR"
    )
])


ws_summary.append([
    "Door Number - N/A",
    total_properties,
    door_na,
    round(
        door_na /
        total_properties *
        100
        if total_properties
        else 0,
        2
    ),
    (
        "PASS"
        if door_na == 0
        else "ERROR"
    )
])


# ------------------------------------------------------------
# PHONE
# ------------------------------------------------------------

ws_summary.append([
    "Phone Number - Duplicate Records",
    total_properties,
    phone_result[
        "duplicate_records"
    ],
    round(
        (
            phone_result[
                "duplicate_records"
            ] /
            total_properties *
            100
        )
        if total_properties
        else 0,
        2
    ),
    (
        "PASS"
        if phone_result[
            "duplicate_records"
        ] == 0
        else "ERROR"
    )
])


ws_summary.append([
    "Phone Number - N/A",
    total_properties,
    phone_na,
    round(
        phone_na /
        total_properties *
        100
        if total_properties
        else 0,
        2
    ),
    (
        "PASS"
        if phone_na == 0
        else "ERROR"
    )
])


# ------------------------------------------------------------
# OVERALL
# ------------------------------------------------------------

ws_summary.append([
    "--------------------------------",
    "",
    "",
    "",
    ""
])


ws_summary.append([
    "Properties With At Least One Error",
    total_properties,
    error_records,
    round(
        overall_error_percentage,
        2
    ),
    (
        "PASS"
        if error_records == 0
        else "ERROR"
    )
])


ws_summary.append([
    "Clean Properties",
    total_properties,
    clean_records,
    round(
        overall_pass_percentage,
        2
    ),
    (
        "PASS"
        if clean_records == total_properties
        else "INFO"
    )
])


# ============================================================
# FINAL WORKBOOK FORMATTING
# ============================================================

for ws in wb.worksheets:

    ws.freeze_panes = "A2"

    style_all_cells(ws)

    resize_columns(ws)


# ============================================================
# SAVE EXCEL
# ============================================================

wb.save(
    output_excel
)


# ============================================================
# FINAL CONSOLE REPORT
# ============================================================

print("\n")
print("=" * 100)
print("ASSESSMENT QC COMPLETED")
print("=" * 100)

print("\nAssessment Feature Class:")
print(assessment_fc)

print("\nTotal Properties:")
print(total_properties)

print("\n")
print("--------------------------------------------")
print("PTIN")
print("--------------------------------------------")

print(
    "Duplicate Values:",
    ptin_result[
        "duplicate_values"
    ]
)

print(
    "Duplicate Records:",
    ptin_result[
        "duplicate_records"
    ]
)

print(
    "Error Percentage:",
    round(
        ptin_error_pct,
        2
    ),
    "%"
)


print("\n")
print("--------------------------------------------")
print("OWNER NAME")
print("--------------------------------------------")

print(
    "Duplicate Values:",
    owner_result[
        "duplicate_values"
    ]
)

print(
    "Duplicate Records:",
    owner_result[
        "duplicate_records"
    ]
)

print(
    "N/A Count:",
    owner_na
)

print(
    "Error Percentage:",
    round(
        owner_error_pct,
        2
    ),
    "%"
)


print("\n")
print("--------------------------------------------")
print("DOOR NUMBER")
print("--------------------------------------------")

print(
    "Duplicate Values:",
    door_result[
        "duplicate_values"
    ]
)

print(
    "Duplicate Records:",
    door_result[
        "duplicate_records"
    ]
)

print(
    "N/A Count:",
    door_na
)

print(
    "Error Percentage:",
    round(
        door_error_pct,
        2
    ),
    "%"
)


print("\n")
print("--------------------------------------------")
print("PHONE NUMBER")
print("--------------------------------------------")

print(
    "Duplicate Values:",
    phone_result[
        "duplicate_values"
    ]
)

print(
    "Duplicate Records:",
    phone_result[
        "duplicate_records"
    ]
)

print(
    "N/A Count:",
    phone_na
)

print(
    "Error Percentage:",
    round(
        phone_error_pct,
        2
    ),
    "%"
)


print("\n")
print("--------------------------------------------")
print("EXISTING USAGE")
print("--------------------------------------------")

for category, count in usage_counter.items():

    pct = (
        count /
        total_properties *
        100
        if total_properties
        else 0
    )

    print(
        "{} : {} ({:.2f}%)".format(
            category,
            count,
            pct
        )
    )


print("\n")
print("--------------------------------------------")
print("OVERALL QC")
print("--------------------------------------------")

print(
    "Properties With Error:",
    error_records
)

print(
    "Clean Properties:",
    clean_records
)

print(
    "Overall Error Percentage:",
    round(
        overall_error_percentage,
        2
    ),
    "%"
)

print(
    "Overall Pass Percentage:",
    round(
        overall_pass_percentage,
        2
    ),
    "%"
)


print("\n")
print("--------------------------------------------")
print("EXCEL OUTPUT")
print("--------------------------------------------")

print(output_excel)


print("\n")
print("=" * 100)
print("COMPLETED")
print("=" * 100)