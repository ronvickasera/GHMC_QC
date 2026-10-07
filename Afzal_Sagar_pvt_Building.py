import arcpy
import os
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================
gdb_path = r"I:\PMU\NEOGEO COLONIES\AFZAL SAGAR\AFZAL SAGAR\AFZAL SAGAR.gdb"

feature_class = "Pvt_Buildings"

fc_path = os.path.join(gdb_path, feature_class)

# Output Excel report
output_excel = os.path.join(
    os.path.dirname(gdb_path),
    "Pvt_Buildings_QC_Report.xlsx"
)


# ============================================================
# CHECK INPUT
# ============================================================

if not arcpy.Exists(fc_path):
    raise Exception(
        f"Feature class does not exist:\n{fc_path}"
    )

print("=" * 70)
print("PVT_BUILDINGS QC REPORT")
print("=" * 70)

print(f"GDB              : {gdb_path}")
print(f"Feature Class    : {feature_class}")
print(f"Output Excel     : {output_excel}")
print()


# ============================================================
# CHECK REQUIRED FIELDS
# ============================================================

fields = [f.name for f in arcpy.ListFields(fc_path)]

required_fields = [
    "OBJECTID",
    "Building_ID",
    "Building_Type"
]

missing_fields = [
    field for field in required_fields
    if field not in fields
]

if missing_fields:
    raise Exception(
        "Missing required field(s): "
        + ", ".join(missing_fields)
    )


# ============================================================
# VARIABLES
# ============================================================

total_buildings = 0

building_id_groups = defaultdict(list)
building_type_counts = defaultdict(int)

detail_records = []


# ============================================================
# READ FEATURE CLASS
# ============================================================

print("Reading Pvt_Buildings...")

with arcpy.da.SearchCursor(
    fc_path,
    [
        "OBJECTID",
        "Building_ID",
        "Building_Type"
    ]
) as cursor:

    for row in cursor:

        objectid = row[0]
        building_id = row[1]
        building_type = row[2]

        total_buildings += 1

        # ----------------------------------------------------
        # BUILDING ID DUPLICATE CHECK
        # ----------------------------------------------------

        # Convert None to readable value
        if building_id is None:
            building_id_key = "<NULL>"
        else:
            building_id_key = str(building_id).strip()

            if building_id_key == "":
                building_id_key = "<EMPTY>"

        building_id_groups[building_id_key].append(
            objectid
        )

        # ----------------------------------------------------
        # BUILDING TYPE CATEGORY
        # ----------------------------------------------------

        if building_type is None:
            building_type_key = "<NULL>"
        else:
            building_type_key = str(building_type).strip()

            if building_type_key == "":
                building_type_key = "<EMPTY>"

        building_type_counts[building_type_key] += 1

        # ----------------------------------------------------
        # DETAIL RECORD
        # ----------------------------------------------------

        detail_records.append(
            (
                objectid,
                building_id,
                building_type
            )
        )


# ============================================================
# IDENTIFY DUPLICATE BUILDING IDs
# ============================================================

duplicate_records = []

duplicate_id_count = 0
duplicate_feature_count = 0

for building_id, objectids in building_id_groups.items():

    if len(objectids) > 1:

        duplicate_id_count += 1
        duplicate_feature_count += len(objectids)

        for oid in objectids:

            # Find Building_Type for this OBJECTID
            record = next(
                (
                    r for r in detail_records
                    if r[0] == oid
                ),
                None
            )

            building_type = record[2] if record else None

            duplicate_records.append(
                (
                    building_id,
                    oid,
                    building_type,
                    len(objectids)
                )
            )


# ============================================================
# PRINT RESULTS TO CONSOLE
# ============================================================

print()
print("-" * 70)
print("SUMMARY")
print("-" * 70)

print(f"Total Pvt_Buildings          : {total_buildings}")
print(f"Unique Building_ID values    : {len(building_id_groups)}")
print(f"Duplicate Building_ID groups : {duplicate_id_count}")
print(f"Features involved in dupes   : {duplicate_feature_count}")

print()
print("-" * 70)
print("BUILDING TYPE CATEGORIES")
print("-" * 70)

for category, count in sorted(
    building_type_counts.items(),
    key=lambda x: x[0]
):

    print(f"{category:<40} : {count}")

print()
print("-" * 70)
print("DUPLICATE BUILDING IDs")
print("-" * 70)

if duplicate_records:

    for record in duplicate_records:

        building_id = record[0]
        oid = record[1]
        building_type = record[2]
        duplicate_count = record[3]

        print(
            f"Building_ID={building_id} | "
            f"OBJECTID={oid} | "
            f"Building_Type={building_type} | "
            f"Duplicate_Count={duplicate_count}"
        )

else:

    print("No duplicate Building_ID values found.")


# ============================================================
# CREATE EXCEL WORKBOOK
# ============================================================

print()
print("Creating Excel report...")

wb = Workbook()


# ============================================================
# STYLES
# ============================================================

header_fill = PatternFill(
    "solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

title_font = Font(
    bold=True,
    size=14
)

subheader_fill = PatternFill(
    "solid",
    fgColor="D9EAF7"
)

thin_border = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin")
)


def format_header(ws, row_number):

    for cell in ws[row_number]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        cell.border = thin_border


def auto_width(ws):

    for column_cells in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            try:

                value_length = len(
                    str(cell.value)
                )

                if value_length > max_length:
                    max_length = value_length

            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 3,
            50
        )


# ============================================================
# SHEET 1 - SUMMARY
# ============================================================

ws = wb.active
ws.title = "Summary"

ws["A1"] = "Pvt_Buildings QC Report"
ws["A1"].font = title_font

ws["A3"] = "Parameter"
ws["B3"] = "Value"

format_header(ws, 3)

summary_data = [
    ("Feature Class", feature_class),
    ("Total Buildings", total_buildings),
    ("Unique Building_ID", len(building_id_groups)),
    ("Duplicate Building_ID Groups", duplicate_id_count),
    ("Features Involved in Duplicates", duplicate_feature_count),
    ("Building_Type Categories", len(building_type_counts)),
    (
        "QC Status",
        "DUPLICATES FOUND"
        if duplicate_id_count > 0
        else "NO DUPLICATES"
    )
]

row_number = 4

for parameter, value in summary_data:

    ws.cell(
        row=row_number,
        column=1,
        value=parameter
    )

    ws.cell(
        row=row_number,
        column=2,
        value=value
    )

    row_number += 1

auto_width(ws)


# ============================================================
# SHEET 2 - BUILDING TYPE COUNT
# ============================================================

ws = wb.create_sheet(
    "Building_Type_Count"
)

ws["A1"] = "Building_Type Category Count"
ws["A1"].font = title_font

headers = [
    "Building_Type",
    "Count",
    "Percentage"
]

for col, header in enumerate(headers, 1):

    ws.cell(
        row=3,
        column=col,
        value=header
    )

format_header(ws, 3)

row_number = 4

for category, count in sorted(
    building_type_counts.items(),
    key=lambda x: x[0]
):

    percentage = (
        count / total_buildings * 100
        if total_buildings > 0
        else 0
    )

    ws.cell(
        row=row_number,
        column=1,
        value=category
    )

    ws.cell(
        row=row_number,
        column=2,
        value=count
    )

    ws.cell(
        row=row_number,
        column=3,
        value=round(percentage, 2)
    )

    row_number += 1

auto_width(ws)


# ============================================================
# SHEET 3 - DUPLICATE BUILDING ID
# ============================================================

ws = wb.create_sheet(
    "Duplicate_Building_ID"
)

ws["A1"] = "Duplicate Building_ID QC"
ws["A1"].font = title_font

headers = [
    "Building_ID",
    "OBJECTID",
    "Building_Type",
    "Duplicate_Count"
]

for col, header in enumerate(headers, 1):

    ws.cell(
        row=3,
        column=col,
        value=header
    )

format_header(ws, 3)

row_number = 4

if duplicate_records:

    for record in duplicate_records:

        for col, value in enumerate(
            record,
            1
        ):

            ws.cell(
                row=row_number,
                column=col,
                value=value
            )

        row_number += 1

else:

    ws.cell(
        row=4,
        column=1,
        value="No duplicate Building_ID found"
    )

auto_width(ws)


# ============================================================
# SHEET 4 - DETAIL
# ============================================================

ws = wb.create_sheet(
    "Pvt_Buildings_Detail"
)

ws["A1"] = "Pvt_Buildings Feature Details"
ws["A1"].font = title_font

headers = [
    "OBJECTID",
    "Building_ID",
    "Building_Type"
]

for col, header in enumerate(headers, 1):

    ws.cell(
        row=3,
        column=col,
        value=header
    )

format_header(ws, 3)

row_number = 4

for record in detail_records:

    for col, value in enumerate(
        record,
        1
    ):

        ws.cell(
            row=row_number,
            column=col,
            value=value
        )

    row_number += 1

auto_width(ws)


# ============================================================
# ADD FILTERS
# ============================================================

for sheet_name in [
    "Building_Type_Count",
    "Duplicate_Building_ID",
    "Pvt_Buildings_Detail"
]:

    ws = wb[sheet_name]

    if ws.max_row >= 3:

        ws.auto_filter.ref = ws.dimensions


# ============================================================
# FREEZE PANES
# ============================================================

for ws in wb.worksheets:

    ws.freeze_panes = "A4"


# ============================================================
# SAVE EXCEL
# ============================================================

wb.save(output_excel)


# ============================================================
# FINAL MESSAGE
# ============================================================

print()
print("=" * 70)
print("QC REPORT GENERATED SUCCESSFULLY")
print("=" * 70)

print(f"Excel Report:")
print(output_excel)

print()
print(f"Total Buildings          : {total_buildings}")
print(f"Duplicate Building_ID    : {duplicate_id_count}")
print(f"Building_Type Categories : {len(building_type_counts)}")

if duplicate_id_count > 0:

    print()
    print("WARNING: Duplicate Building_ID values were found.")

else:

    print()
    print("PASS: No duplicate Building_ID values found.")

print("=" * 70)